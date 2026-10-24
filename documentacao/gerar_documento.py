"""Gera a documentação técnica; não participa da execução do dashboard.

Usa o ambiente de documentos da ferramenta, separado das dependências da aplicação.
"""
from pathlib import Path
import sys
from math import atan2, cos, sin
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
from arquitetura_projeto import (INTRODUCAO, REQUISICOES, ETL, FUNCIONAIS,
                                NAO_FUNCIONAIS, FUTURO, LIMITE)
from gerar_arquitetura import gerar
arquitetura_imagem = gerar()
OUT = ROOT / 'Arquitetura_do_dashboard_Heliopolis.docx'
QA = ROOT / 'verificacao'
QA.mkdir(exist_ok=True)

# O fluxograma é desenhado em alta resolução para permanecer legível no Word.
im = Image.new('RGB', (1800, 2120), 'white')
draw = ImageDraw.Draw(im)
font_path = 'C:/Windows/Fonts/arial.ttf'
font = ImageFont.truetype(font_path, 31)
small = ImageFont.truetype(font_path, 26)
bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 33)
def box(rect, title, body, fill='#EDF3F7'):
    draw.rounded_rectangle(rect, radius=15, fill=fill, outline='#547084', width=3)
    x0,y0,x1,y1=rect
    lines=[(title,bold)]+[(line,font) for line in body.split('\n')]
    y=(y0+y1)/2-((len(lines)-1)*40+34)/2
    for line,f in lines:
        draw.text(((x0+x1)/2,y),line,font=f,fill='#142A39',anchor='mt')
        y+=40
def arrow(points, label=None, labelpos=None):
    draw.line(points,fill='#435C6B',width=5)
    x,y=points[-1]; px,py=points[-2]; a=atan2(y-py,x-px)
    draw.polygon([(x,y),(x-20*cos(a-.45),y-20*sin(a-.45)),(x-20*cos(a+.45),y-20*sin(a+.45))],fill='#435C6B')
    if label: draw.text(labelpos,label,font=small,fill='#243E50',anchor='mm')

# Há uma única origem meteorológica; os dados de saúde chegam por uma tabela local à direita.
box((40,40,1160,220),'Série ERA5 pelo Open-Meteo','Sessão existente → JSON local → API se não houver arquivo\nA consulta HTTPS GET retorna uma resposta JSON')
box((40,310,1160,510),'dados.py','Validar campos e unidades • converter horários para UTC\nConstruir a grade horária • remover duplicatas\nPreservar lacunas e registrar a qualidade')
arrow([(600,220),(600,310)])
box((40,600,1160,780),'painel.py','Recortar 2021 a 2025 e aplicar o filtro de ano\nDefinir o limiar analítico de temperatura')
arrow([(600,510),(600,600)])
box((40,870,1160,1070),'analise.py','Calcular dias e noites com cobertura mínima de 90%\nIdentificar sequências de pelo menos três dias\nResumir intensidade e frequência do calor')
arrow([(600,780),(600,870)])
box((40,1160,1160,1360),'saude.py','Dias quentes x esperado do mesmo ano, mês e dia da semana\nIntervalo de 95% por reamostragem • mínimo de 10 dias\nComparação também por faixa de temperatura')
arrow([(600,1070),(600,1160)])
box((1260,1160,1760,1360),'Tabela de internações','SIH/SUS • cidade de São Paulo\nContagens diárias, 2021–2025\nPreparada fora do dashboard',fill='#F3EEDF')
arrow([(1260,1260),(1160,1260)])
box((40,1450,1160,1650),'Streamlit em uma única página','Gráficos de temperatura, frequência, horário, noite e umidade\nCartões, tabela de sequências e seção Calor e internações\nAlterar os filtros recalcula as análises',fill='#E4EEE8')
arrow([(600,1360),(600,1450)])
draw.text((600,1750),'Reutilização da série',font=bold,fill='black',anchor='mt')
draw.text((600,1810),'Sessão e cache evitam repetir a coleta a cada interação.',font=font,fill='black',anchor='mt')
draw.text((600,1900),'Alcance dos resultados',font=bold,fill='black',anchor='mt')
draw.text((600,1960),'Contexto regional estimado, sem medição dentro das moradias.',font=font,fill='black',anchor='mt')
draw.text((600,2010),'Internações são da cidade inteira e só do SUS, não do bairro.',font=font,fill='black',anchor='mt')
draw.text((600,2060),'Qualidade de vida e dados do território ainda não integrados.',font=font,fill='black',anchor='mt')
flow = QA / 'fluxograma.png'; im.save(flow)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2); sec.right_margin=Cm(2)
sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)
for name,size in [('Normal',10.5),('Title',25),('Heading 1',17),('Heading 2',12),('Caption',9)]:
    style=doc.styles[name]; style.font.name='Calibri'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.12
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Normal'].paragraph_format.widow_control=True
# O modelo padrão pode herdar uma borda azul no estilo Title.
# Removemos bordas de parágrafo dos estilos para manter o título limpo.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
head=sec.header.paragraphs[0]; head.text='HELIÓPOLIS  |  DOCUMENTAÇÃO DA ARQUITETURA'
head.runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Versão de 24 outubro 2026  •  Página ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
for r in foot.runs: r.font.size=Pt(8)
doc.core_properties.title='Arquitetura e metodologia do dashboard de Heliópolis'
doc.core_properties.subject='Fluxo dos dados, ERA5, cálculos e responsabilidades dos módulos'
doc.core_properties.author='Projeto Heliópolis'
def p(t): return doc.add_paragraph(t)
def h(t): return doc.add_heading(t,level=2)
def page(t): doc.add_page_break(); doc.add_heading(t,level=1)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'24475C' if ri==0 else ('F0F4F7' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.05
                for r in para.runs:
                    r.font.size=Pt(9.5)
                    if ri==0: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    p('')
    return t
def formula(text):
    # OMML mantém as expressões editáveis no Word, sem depender de imagens.
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath'); run=OxmlElement('m:r'); txt=OxmlElement('m:t'); txt.text=text
    run.append(txt); math.append(run); para._p.append(math)

doc.add_paragraph('Arquitetura e metodologia do dashboard de Heliópolis','Title')
p('Guia de estudo do fluxo de dados e dos cálculos implementados')
p('Este documento explica a implementação atual do dashboard: como os dados de clima e de internações são obtidos, tratados e utilizados; quais cálculos produzem cada gráfico; e como os arquivos Python se comunicam. O período analisado é 2021–2025, definido como os cinco anos completos anteriores a 2026.')
h('Engenharia de software e arquitetura')
p(INTRODUCAO)
h('Requisitos funcionais implementados')
table(['Código','Requisito','Comportamento'], FUNCIONAIS, [1.6,3.7,11.7])
page('Requisitos não funcionais e limites operacionais')
p('Os requisitos funcionais descrevem ações observáveis, como filtrar anos e calcular episódios. Os requisitos não funcionais definem condições de qualidade e operação dessas ações. A tabela registra os critérios adotados no código, com os limites que ainda não foram avaliados ou implementados.')
table(['Código','Atributo','Critério atual'], NAO_FUNCIONAIS, [1.6,3.7,11.7])
h('Escopo ainda não implementado')
p(FUTURO)
p(LIMITE)
h('Como verificar os requisitos')
p('Os testes automatizados verificam filtros, cobertura, lacunas, fronteiras de datas e tratamento da resposta da API. A inspeção dos metadados permite rastrear a origem dos dados. Sessão e cache reduzem trabalho repetido, mas não constituem teste de carga ou garantia de tempo de resposta. Falhas previstas recebem tratamento; indisponibilidade externa, falhas inesperadas e alterações no formato das fontes ainda podem interromper a análise.')

page('Arquitetura das requisições e do processamento')
p('A aplicação consome a API REST do Open-Meteo ou lê o JSON dessa mesma fonte salvo localmente. Os dados de saúde entram por uma tabela local, gerada por um pipeline offline que roda no terminal. As setas representam as trocas de dados e a sequência das responsabilidades; não representam microsserviços independentes.')
doc.add_picture(str(arquitetura_imagem),width=Cm(17))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.inline_shapes[-1]._inline.docPr.set('descr','API externa recebe GET e retorna JSON ao módulo de extração. Arquivos locais também alimentam a extração. Transformação com Pandas e carga em memória alimentam as análises de calor e o módulo de internações, que lê uma tabela local gerada por um pipeline offline com dados do DATASUS. O Streamlit troca interface e filtros com o navegador.')
doc.add_paragraph('Figura de arquitetura  Requisições externas, ETL em memória, tabela de internações e pipeline offline de saúde','Caption')
p(REQUISICOES)
p(ETL)
p(LIMITE)

page('Contexto do estudo e explicação do ERA5')
h('Objetivo e alcance atual')
p('O projeto utiliza Ciência de Dados para estudar as condições de calor relacionadas à saúde e à qualidade de vida dos moradores de Heliópolis. A etapa atual combina a série ERA5 disponibilizada pelo Open-Meteo com uma tabela de internações hospitalares do SIH/SUS da cidade de São Paulo; a investigação de associações com qualidade de vida e com o território de Heliópolis depende de dados compatíveis ainda não obtidos.')
p('Como não há medições dentro da comunidade e os dados de saúde cobrem a cidade inteira (apenas internações pelo SUS), a versão documentada descreve condições regionais de calor e as compara com as internações da cidade. Ela não mede temperaturas dentro das moradias, sintomas, qualidade do sono ou impactos individuais.')
h('O que é ERA5')
p('ERA5 é a quinta geração de reanálise atmosférica do ECMWF, produzida para o Serviço de Mudanças Climáticas Copernicus. Uma reanálise combina observações históricas, como medições de estações e satélites, com um modelo físico e técnicas de assimilação de dados para estimar o estado passado da atmosfera. Oferece séries horárias globais desde 1940. [1, 2]')
p('Isso permite contextualizar as condições meteorológicas de uma região ao longo do tempo. Os valores são estimativas de reanálise, e não leituras de um termômetro instalado no ponto consultado. A resolução nativa do ERA5 é de aproximadamente 31 km; produtos distribuídos em grade regular usam 0,25°. O Open-Meteo descreve essa grade como cerca de 25 km. [1, 2, 3]')
p('O projeto acessa ERA5 pelo Open-Meteo com models=era5. Open-Meteo é a interface de acesso; ERA5 é a base meteorológica. A API pode aplicar ajustes estatísticos de altitude por padrão: o código não define elevation=nan para desativá-los. Isso não transforma a série em uma medição do microclima de Heliópolis. [3]')
p('A série regional ajuda a explorar intensidade, frequência e persistência do calor. Para relacioná-la à experiência dos moradores de Heliópolis, ainda serão necessários dados de saúde e qualidade de vida compatíveis no tempo e no território do bairro.')

page('1 Fluxograma da sequência dos dados')
p('O fluxo usa uma única fonte meteorológica: ERA5 pelo Open-Meteo. A mesma série pode ser obtida pela API ou por seu JSON salvo. Depois da validação e do recorte, ela alimenta todos os gráficos e indicadores. O módulo saude.py cruza os dias quentes com uma tabela local de internações, que não vem da internet em tempo de execução.')
doc.add_picture(str(flow),width=Cm(17))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.inline_shapes[-1]._inline.docPr.set('descr','Fluxograma da série ERA5 acessada pelo Open-Meteo, por sessão, JSON local ou API. Após validação, recorte e cálculo, os indicadores seguem para o módulo de saúde, que também lê a tabela local de internações, e são apresentados no Streamlit.')
doc.add_paragraph('Figura 1  Sequência implementada de entrada, tratamento, análise, comparação com internações e apresentação dos dados','Caption')
p('Os módulos se comunicam por chamadas de funções e objetos Python em memória. Não há banco de dados, API interna ou serviço de processamento separado.')

page('2 Coleta e tratamento da série ERA5')
h('Prioridade de carregamento')
p('Ao abrir a aplicação, iniciar_base verifica primeiro se existe um conjunto em st.session_state. Se existir, ele é reutilizado. Na primeira carga, procura dados_locais/heliopolis_open_meteo.json. Se o arquivo existir, ele é lido e validado. Somente se não existir, carregar_online consulta o Open-Meteo. Um arquivo inválido gera indisponibilidade; não há substituição automática por dados artificiais.')
p('A consulta usa urllib, da biblioteca padrão, com o endpoint https://archive-api.open-meteo.com/v1/archive. Solicita latitude −23,615 e longitude −46,590, modelo ERA5, timestamps Unix e fuso UTC. O ponto é aproximado e não corresponde a um centroide territorial oficialmente delimitado pelo projeto.')
table(['Variável solicitada','Nome interno','Unidade'],[
('Temperatura do ar a 2 m','temperatura_c','°C'),('Umidade relativa a 2 m','umidade_pct','%'),
('Vento a 10 m','vento_ms','m/s'),('Radiação solar','radiacao_wm2','W/m²'),('Precipitação','precipitacao_mm','mm')],[6.5,7.5,3])
p('A resposta JSON contém listas horárias e unidades. interpretar_api verifica os campos, converte timestamps e cria um DataFrame indexado por horário. A interface atual utiliza principalmente temperatura e umidade; vento, radiação e precipitação permanecem no processamento auxiliar.')
h('Limpeza e cobertura')
p('limpar_horarios ordena a série, remove horários duplicados mantendo a primeira ocorrência e reindexa em uma grade com todas as horas esperadas. Em 2021–2025 são 43.824 horas e 1.826 dias. Registros ausentes e valores inválidos viram NaN; não são interpolados nem substituídos por zero.')
table(['Variável','Intervalo de consistência aceito'],[('Temperatura','−90 a 65 °C'),('Umidade','0 a 100%'),('Vento','0 a 150 m/s'),('Radiação','0 a 1.600 W/m²'),('Precipitação horária','0 a 1.000 mm')],[6.5,10.5])
p('Esses limites amplos verificam consistência básica. Não são limites de saúde e não substituem controle completo de qualidade. Para cada variável, são registrados horas esperadas, válidas, ausentes ou inválidas, valores fora dos limites e percentual de cobertura.')
p('montar_conjunto reúne horários, qualidade, número de duplicatas, metadados e resposta de origem. Os metadados preservam modelo, período, coordenadas, data de coleta, fuso e URL da consulta, permitindo rastrear a proveniência.')

page('3 Cálculos diários e gráficos de temperatura')
h('Regra de cobertura diária')
p('Se n é a quantidade de registros válidos e H a quantidade de horas esperadas no dia, a cobertura é calculada pela expressão abaixo. O dia é válido com pelo menos 90% das horas de temperatura.')
formula('Cobertura diária (%) = 100 × n / H')
p('Em um dia de 24 horas, são necessários pelo menos 22 registros válidos. Um dia inválido fica com temperatura agregada e classificação de calor ausentes; ele não é interpretado como um dia ameno. Os agrupamentos usam o fuso America/Sao_Paulo.')
h('Temperatura ao longo do tempo')
p('Para cada dia válido, a média é a soma das temperaturas válidas dividida pela quantidade de registros; a máxima e a mínima são os extremos dessas amostras horárias.')
formula('Temperatura média diária = soma das temperaturas válidas / n')
p('Até 370 dias, o gráfico exibe medidas diárias. Em períodos maiores, exibe a média mensal de cada medida diária. A linha mensal de máxima é a média das máximas dos dias válidos daquele mês, não o maior valor ocorrido no mês. O mesmo vale para média e mínima. As agregações mensais não impõem um segundo limiar mínimo de dias válidos.')
p('Essa visualização permite explorar intensidade e sazonalidade. Máximas horárias podem não captar picos ocorridos entre as amostras. Cartões e episódios continuam usando medidas diárias, sem a suavização mensal do gráfico.')
h('Frequência de dias quentes')
p('Um dia válido é classificado como quente quando sua máxima é maior ou igual ao limiar L. O padrão é 30 °C, ajustável de 25 a 40 °C em passos de 0,5 °C. Esse é um critério analítico, não um limite clínico.')
formula('Frequência anual (%) = 100 × dias válidos com máxima ≥ L / dias válidos')
p('Exemplo: 70 dias acima do limiar entre 350 dias válidos representam 20%. O denominador não é automaticamente 365 ou 366. Não há um limiar adicional de cobertura anual; a frequência descreve os dias disponíveis e não estabelece tendência climática com apenas cinco anos.')
h('Calor por horário')
p('A série horária é convertida ao horário local e agrupada de 0h a 23h. Para cada hora, calcula-se a média dos registros válidos daquela hora ao longo do período. O ponto das 14h representa a média de todos os registros válidos às 14h.')
p('Esse gráfico inclui registros válidos de dias que eventualmente não atingem 90% de cobertura diária. Não é filtrado pelo limiar L. Ele descreve o perfil térmico médio ao longo da rotina, sem representar exposição pessoal medida.')

page('4 Noites umidade e persistência do calor')
h('Temperatura à noite')
p('Cada noite começa às 22h e termina às 05h59 do dia seguinte. Os registros da madrugada são associados à data de início da noite. A janela completa deve estar dentro do período consultado e ter pelo menos 90% de cobertura.')
p('Uma noite usual tem oito registros horários. Como sete correspondem a 87,5%, são necessários os oito para atingir o critério. O código também considera janelas de sete ou nove horas em transições históricas de horário de verão.')
formula('Média noturna = soma das temperaturas noturnas válidas / registros válidos')
p('Se a tabela tem mais de 370 linhas, a visualização apresenta a média mensal das médias noturnas, agrupando pela data de início da noite. Caso contrário, mostra cada noite. Noites parciais nas bordas são excluídas. A análise descreve o ambiente durante o descanso, sem medir sono ou temperatura dentro das casas.')
h('Calor e umidade')
p('Cada ponto da dispersão corresponde às médias de temperatura e umidade de um dia. O dia precisa ser válido para temperatura e ter cobertura de pelo menos 90% para umidade. A verificação é independente por variável: as horas efetivamente usadas nas duas médias podem diferir.')
p('O eixo horizontal recebe a temperatura média diária; o vertical recebe a umidade média diária. Todos os dias elegíveis aparecem, inclusive aqueles abaixo do limiar. Não há cálculo de Pearson, regressão, sensação térmica, UTCI ou índice fisiológico nessa visualização.')
h('Sequências de calor persistente')
p('episodios_quentes percorre os dias em ordem cronológica e identifica sequências de pelo menos três dias consecutivos com máxima ≥ L. Um dia abaixo do limiar, inválido ou ausente interrompe a sequência. Um episódio longo conta uma vez, e não como várias janelas sobrepostas de três dias.')
table(['Medida do episódio','Como é calculada'],[('Início e fim','Primeiro e último dia da sequência'),('Duração','Número de dias consecutivos'),('Maior temperatura','Maior máxima diária da sequência'),('Média das máximas','Média das máximas diárias da sequência'),('Borda do recorte','Indica contato com o início ou o fim da consulta')],[5.5,11.5])
p('Ao filtrar um ano, episódios que atravessam a virada podem ser cortados. A marcação de borda indica que podem continuar fora do recorte. A definição de três dias é exploratória e não corresponde automaticamente a uma onda de calor oficial.')

page('5 Cartões e análises auxiliares')
table(['Cartão','Cálculo implementado'],[('Maior temperatura horária','Maior máxima diária entre os dias válidos'),('Dias com máxima ≥ limiar','Soma dos dias válidos classificados como quentes'),('Sequências de calor','Quantidade de episódios com três dias ou mais'),('Temperatura média noturna','Média das médias de todas as noites válidas')],[6,11])
p('No cartão noturno, cada noite válida tem o mesmo peso. Quando não há suporte suficiente, resultados ausentes são apresentados como travessão, sem confundir ausência com zero.')
h('Cálculos auxiliares fora da tela atual')
p('O módulo também calcula horas observadas acima do limiar, médias diárias de vento e radiação e chuva diária. A chuva exige 100% de cobertura antes de somar os valores; as demais médias auxiliares exigem 90%. Nenhuma lacuna é extrapolada. Esses resultados podem ser consultados nos arquivos gerados pelo terminal.')

page('6 Calor e internações (SIH/SUS)')
h('De onde vêm os dados de saúde')
p('As internações vêm do Sistema de Informações Hospitalares do SUS (SIH/SUS), arquivos AIH Reduzida (RD) de São Paulo, publicados pelo DATASUS. O dashboard não consulta o DATASUS: lê a tabela pronta dados_saude/internacoes_diarias_sp.csv, que foi gerada antes por um pipeline separado. Um arquivo .json ao lado dela registra a fonte, os filtros, a data de geração e o número de linhas.')
table(['Coluna da tabela','Conteúdo'],[('data','Dia da internação, de 01/01/2021 a 31/12/2025'),('grupo_cid','I00-I99 (aparelho circulatório) ou J00-J99 (aparelho respiratório)'),('faixa_etaria','0-14, 15-64, 65+ ou sem idade'),('n_internacoes','Número de internações naquele dia, grupo e faixa; dias sem internação aparecem com zero')],[4,13])
p('Entram apenas moradores da cidade de São Paulo (código IBGE 355030) com AIH normal. A tabela não contém CEP, idade exata, hospital nem qualquer dado individual. Ela descreve a cidade inteira e somente as internações pelo SUS; não identifica moradores de Heliópolis.')
h('Pipeline offline em exploracao_saude/')
p('Os scripts da pasta exploracao_saude/ rodam no terminal, com ambiente próprio (Python 3.12 e PySUS), e não fazem parte da aplicação Streamlit. O catálogo do PySUS tem 52 dos 60 meses de 2021–2025; os 8 meses ausentes foram obtidos no FTP do DATASUS, em arquivos .dbc. Os arquivos de competências até 2026-08 foram usados para completar as internações de 2025, que chegam com atraso.')
table(['Script','O que faz'],[('01 e 01b','Listam os arquivos disponíveis no catálogo do PySUS e no FTP'),('02 e 03','Exploram um mês: colunas, perfil e contagem diária'),('04 e 05','Baixam mês a mês, reduzem cada arquivo a contagens, apagam o arquivo bruto e montam as contagens mensal e diária'),('06 e 08','Cruzam as internações com a temperatura e conferem o método do painel com regressão de Poisson, em vários limiares'),('07','Exporta a tabela dados_saude/internacoes_diarias_sp.csv que o dashboard lê')],[3,14])
p('O relatório exploracao_saude/RELATORIO.md resume o que foi feito e conclui que os dados do SIH permitem uma série diária da cidade, mas não chegam ao bairro de Heliópolis.')
h('Como a comparação é calculada')
p('Dias de calor e dias comuns diferem também por mês do ano, ano e dia da semana, que afetam as internações independentemente do clima. Para reduzir esse efeito, cada dia quente é comparado com o valor esperado: a média das internações dos dias não quentes do mesmo ano, mês e dia da semana. Os dias quentes ficam fora dessa média. Um dia quente só entra se houver pelo menos dois dias de referência.')
formula('Diferença (%) = 100 × (internações reais / internações esperadas nos dias quentes − 1)')
p('O intervalo de 95% vem de sortear, com reposição, os dias quentes 2.000 vezes (reamostragem). Se o intervalo inclui o zero, o painel não mostra diferença clara. Com menos de 10 dias quentes ou menos de 10 dias de comparação, o resultado não é calculado e a tela orienta o usuário.')
table(['Gráfico da seção','O que compara'],[('Internações nos dias quentes','Dias com máxima ≥ limiar do painel, por grupo de doença e por idade (todas as idades e 65 anos ou mais)'),('Internações por faixa de temperatura','Cada faixa de temperatura máxima (abaixo de 20, 20 a 24, 24 a 27, 27 a 30 e 30 °C ou mais) contra os dias fora da faixa, para todas as idades; não depende do limiar')],[5.5,11.5])
p('Os gráficos de barras desta seção usam Altair, instalado junto com o Streamlit, para deixar os nomes do eixo na horizontal. saude.py faz as contas e interpretacao.py escreve as leituras em linguagem simples (leitura_saude e leitura_faixas), sem desenhar gráficos.')
h('Limites de interpretação')
p('É uma comparação descritiva e mostra associação, não causa. Os dados são da cidade inteira e só do SUS, enquanto o calor é uma estimativa regional. Diferenças pequenas podem ser acaso, e faixas muito frias ou muito quentes têm menos dias e resultados mais incertos. O texto da tela traz esses avisos junto de cada resultado.')

page('7 Responsabilidades dos arquivos principais')
h('app.py e painel.py')
p('app.py é o ponto de entrada do comando streamlit run app.py. Ele importa main de painel.py e executa a função. Não contém regras de cálculo ou coleta.')
p('painel.py coordena a aplicação. iniciar_base obtém a base; recortar_base aplica o ano selecionado; main chama indicadores_diarios, indicadores_noturnos e episodios_quentes. mostrar_exposicao desenha os cartões e gráficos de evolução e frequência; mostrar_rotina apresenta horário, noite e umidade; mostrar_persistencia exibe os episódios; mostrar_saude desenha a seção Calor e internações, com barras_agrupadas (Altair) para os dois gráficos de barras.')
p('A interface chama numero de apoio.py para formatar os valores apresentados. Parte das agregações visuais, como médias mensais e perfil horário, ainda está em painel.py. A separação entre apresentação e cálculo é, portanto, parcial.')
h('dados.py')
p('Centraliza constantes do estudo, validação de períodos, consulta HTTP, leitura e exportação JSON, grade horária e limpeza da série ERA5. Não depende do Streamlit. ErroDados fornece mensagens que podem ser tratadas pela interface ou pelo terminal.')
table(['Chave do conjunto','Conteúdo'],[('horarios','DataFrame com a série horária tratada'),('qualidade','DataFrame com cobertura e inconsistências'),('duplicatas','Número de horários duplicados removidos'),('metadados','Dicionário com origem e parâmetros da coleta'),('resposta_api','Dicionário com a resposta de origem do período')],[5,12])
h('analise.py')
p('Recebe tabelas do Pandas e devolve resultados analíticos sem exigir interface. indicadores_diarios agrega dias e verifica cobertura; indicadores_noturnos organiza janelas de descanso; episodios_quentes identifica sequências; resumo_anual calcula contagens, médias e frequências. gerar_relatorio monta o texto auxiliar com critérios, indicadores e limitações.')
h('saude.py e interpretacao.py')
p('saude.py lê a tabela de internações (carregar_internacoes), escolhe grupo de doença e idade (serie_diaria) e compara os dias quentes e cada faixa de temperatura com o esperado (comparar_todos e comparar_por_faixas). Recebe a tabela diária de analise.py e não desenha gráficos nem escreve textos.')
p('interpretacao.py reúne os textos explicativos do dashboard: fontes, glossário, conclusões e as leituras automáticas que citam os números do recorte, como leitura_saude e leitura_faixas. Não faz contas estatísticas.')
page('8 Comunicação execução e persistência')
h('apoio.py e arquitetura_projeto.py')
p('apoio.py contém apenas numero, responsável por casas decimais, separadores em português e ausência exibida como travessão.')
p('arquitetura_projeto.py centraliza o texto dos requisitos, a descrição arquitetural e o caminho da imagem usados no dashboard e neste documento. assets/arquitetura_projeto.png é a imagem compartilhada. Os geradores em documentacao/ preparam os artefatos e não são executados pelo dashboard; suas dependências de edição de documentos não fazem parte da aplicação Streamlit.')
h('gerar_relatorio.py')
p('É uma entrada alternativa pelo terminal. Usa dados.py e analise.py para gerar um JSON de origem, CSVs horário, diário, noturno, episódios, resumo anual e qualidade, além de parametros.json e relatorio_inicial.txt. O argumento --arquivo permite repetir a análise offline; --inicio e --fim delimitam o período; --saida define o destino.')
p('Seus critérios são fixos em 30 °C, três dias e 90% de cobertura. O script não recebe automaticamente o limiar selecionado na sessão do dashboard. A execução pode substituir os arquivos de saída com os mesmos nomes.')
h('Reexecução e cache')
p('Ao alterar um filtro, o Streamlit executa novamente a página. st.session_state conserva o conjunto ERA5 da sessão. carregar_online usa cache de 24 horas e até seis entradas.')
p('O ano altera todas as análises. O limiar altera a classificação dos dias, frequência anual e episódios. Ele também define quais dias entram na comparação com internações nos dias quentes, mas não altera as temperaturas, os gráficos de evolução, perfil horário, noite e dispersão, nem o gráfico por faixa de temperatura.')
p('A consulta online do painel não grava automaticamente o JSON em disco. Essa persistência é feita pelo script de relatório. Um JSON alterado no disco também não substitui por si só o conjunto já carregado na sessão. Uma falha na primeira carga não provoca tentativas repetidas a cada filtro na mesma sessão.')
h('Objetos que atravessam os módulos')
p('A comunicação usa funções importadas, DataFrames, dicionários e datas. Não existe uma API interna. O caminho principal é app → painel → dados e analise; para a seção de saúde, painel chama saude e interpretacao, e saude recebe a tabela diária de analise. painel também importa textos de arquitetura_projeto e a formatação de apoio. gerar_relatorio chama dados e analise sem iniciar Streamlit.')

page('9 Armazenamento testes e limites de interpretação')
table(['Arquivo ou pasta','Função'],[('dados_locais/heliopolis_open_meteo.json','Base ERA5 de entrada preferencial do painel'),('Outros CSVs e TXT em dados_locais/','Resultados auxiliares gerados pelo terminal'),('dados_saude/','Tabela diária de internações do SIH/SUS (CSV) e arquivo .json com a origem; é lida pelo painel'),('exploracao_saude/','Scripts, saídas e relatório do pipeline offline que gera a tabela de saúde; fora do dashboard'),('requirements.txt','Dependências diretas da aplicação: Streamlit e Pandas'),('.gitignore','Exclui dados locais, ambiente virtual, caches e segredos do Git'),('.venv/ e __pycache__/','Ambiente de execução e cache de bytecode Python')],[7,10])
p('O dashboard lê o JSON ERA5 ou consulta a API quando ele não existe, recalculando os resultados. Os CSVs diários exportados pelo terminal são saídas auxiliares e não alimentam o painel; a tabela de internações em dados_saude/ é a exceção, porque o painel a lê. Imagens do dashboard servem para conferência visual e não alimentam a análise.')
h('Testes automatizados')
table(['Arquivo em tests/','Responsabilidade'],[('test_relatorio.py','Limpeza, cobertura, datas, noites, episódios, API e relatório'),('test_interface.py','Abertura, filtros e ausência de controles técnicos removidos'),('test_saude.py','Comparação com o esperado, intervalo, dias inválidos, faixas de temperatura, tabela sem dados pessoais e seção no painel'),('test_interpretacao.py','Textos dinâmicos, perfil horário, correlação, glossário e conclusões')],[7,10])
p('Os testes utilizam dados sintéticos exclusivamente para validar o comportamento do código. Eles não são inseridos na base apresentada aos moradores ou pesquisadores.')
h('O que os resultados permitem afirmar')
p('A implementação permite descrever o contexto regional de calor e levantar perguntas sobre intensidade, frequência, persistência e condições noturnas. Os critérios de cobertura e a preservação de lacunas tornam explícito o suporte de dados de cada resultado.')
p('Ela não calcula efeitos clínicos, causalidade, desconforto individual, ilha de calor dentro de Heliópolis ou tendências climáticas robustas com cinco anos. O limiar de temperatura não é uma classificação nacional universal de conforto. A grade regional de reanálise não permite quantificar diretamente o excesso térmico dentro do bairro.')
p('A seção de internações mostra associação entre calor regional e internações da cidade inteira pelo SUS; não identifica efeito no bairro nem causa. Para cumprir integralmente o objetivo de impactos na saúde e na qualidade de vida em Heliópolis, será necessário integrar desfechos compatíveis com a exposição ambiental estimada no tempo e no território do bairro, com definições e cobertura conhecidas. Essa etapa exige novos dados e regras analíticas, não apenas novos gráficos.')

page('10 Fontes e rastreabilidade da documentação')
p('A descrição dos módulos e cálculos foi conferida no código local vigente em 24 de outubro de 2026. As referências externas abaixo fundamentam a explicação do ERA5 e de sua disponibilização pelo Open-Meteo. Elas não substituem a verificação dos parâmetros efetivamente usados pelo projeto.')
h('Referências externas')
for text in [
    '[1] ECMWF. ECMWF Reanalysis v5 ERA5. Descrição do produto, séries horárias, origem e resolução nativa. https://www.ecmwf.int/en/forecasts/dataset/ecmwf-reanalysis-v5',
    '[2] Copernicus Climate Change Service. ERA5 hourly data on single levels from 1940 to present. Catálogo do conjunto de reanálise. https://cds.climate.copernicus.eu/datasets/reanalysis-era5-single-levels?tab=overview',
    '[3] Open-Meteo. Historical Weather API. Modelo ERA5, grade de 0,25°, variáveis, unidades, seleção de célula e ajuste estatístico por altitude. https://open-meteo.com/en/docs/historical-weather-api']:
    p(text)
p('Data de consulta das referências de ERA5 e Open-Meteo: 18 de outubro de 2026. A aplicação usa o Open-Meteo como intermediário; não faz download direto do Climate Data Store.')
h('Pontos do código usados como referência')
table(['Arquivo','Funções ou trechos centrais'],[('app.py','Inicialização de main'),('painel.py','iniciar_base, recortar_base, mostrar_exposicao, mostrar_rotina, mostrar_persistencia, mostrar_saude, barras_agrupadas, main'),('dados.py','grade_horaria, limpar_horarios, interpretar_api, montar_conjunto, recortar_periodo, buscar_open_meteo'),('analise.py','indicadores_diarios, indicadores_noturnos, episodios_quentes, resumo_anual'),('saude.py','carregar_internacoes, serie_diaria, comparar_dias_quentes, comparar_por_faixas, comparar_todos'),('interpretacao.py','leitura_saude, leitura_faixas, conclusoes e os textos fixos da página'),('apoio.py e gerar_relatorio.py','numero e entrada alternativa de geração de relatório')],[5.2,11.8])
p('Os nomes acima são relativos à raiz do projeto. Esta documentação registra o estado atual; mudanças futuras em filtros, cobertura, limiares, fontes ou funções devem ser refletidas no documento para preservar a reprodutibilidade.')

doc.save(OUT)
print(OUT)
